import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx', read_only=True)

# Let's check '11. Herbal Products' and '17. Herbal Medicine'
h_prod = wb['11. Herbal Products']
h_med = wb['17. Herbal Medicine']

p_rows = list(h_prod.iter_rows(values_only=True))
m_rows = list(h_med.iter_rows(values_only=True))

print(f"11. Herbal Products: total rows = {len(p_rows)}, non-empty = {len([r for r in p_rows if any(x is not None for x in r)])}")
print(f"17. Herbal Medicine: total rows = {len(m_rows)}, non-empty = {len([r for r in m_rows if any(x is not None for x in r)])}")

# Inspect first few rows of 17. Herbal Medicine
print("\nSample rows from 17. Herbal Medicine:")
for i in range(2, min(len(m_rows), 7)):
    r = m_rows[i]
    print(f"Row {i+1}: Q{r[0]} | {str(r[1])[:50]} | C1: {str(r[3])[:20]} | Ans: {r[8]} | Subtopic: {r[11]}")
