import openpyxl

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx')
ws = wb['9. Neurologic']

for r in range(3, ws.max_row + 1):
    q = str(ws.cell(r, 2).value or '')
    if 'ergotamine และ caffeine 100 mg' in q:
        ws.cell(r, 9).value = 5
        exp_lines = [
            "✅ คำตอบที่ถูกต้อง: ข้อ จ. ไม่มีข้อใดผิด",
            "",
            "💡 Background:",
            "ยาผสม Ergotamine tartrate 1 mg + Caffeine 100 mg (เช่น Cafergot, Avamigran) เป็นยาจำเพาะสำหรับรักษาอาการปวดศีรษะไมเกรนเฉียบพลันระดับปานกลางถึงรุนแรง โดยมีข้อกำหนดด้านขนาดยาและการเฝ้าระวังความปลอดภัยอย่างเคร่งครัด",
            "",
            "🎯 ทำไมข้อนี้ถึงถูก:",
            "ข้อความในข้อ ก, ข, ค, และ ง เป็นคำแนะนำที่ถูกต้องตามหลักเภสัชกรรมคลินิกและฉลากยาทั้งหมด จึงต้องตอบข้อ จ. ไม่มีข้อใดผิด:",
            "• ข้อ ก. ถูกต้อง: ขนาดยาเริ่มต้นแนะนำให้รับประทาน 1-2 เม็ดทันทีที่เริ่มมีอาการปวดศีรษะ (onset of migraine attack)",
            "• ข้อ ข. ถูกต้อง: หากอาการปวดไม่ทุเลา สามารถรับประทานซ้ำได้ครั้งละ 1 เม็ด ทุก 30 นาที",
            "• ข้อ ค. ถูกต้อง: อาการชา เย็น หรือปวดเกร็งที่ปลายมือปลายเท้า (Paresthesia/Cold extremities) เป็นอาการเตือนของภาวะหลอดเลือดหดเกร็งอย่างรุนแรง (Ergotism/Peripheral ischemia) ต้องสั่งให้ผู้ป่วยหยุดยาทันทีและมาพบแพทย์",
            "• ข้อ ง. ถูกต้อง: ขนาดรับประทานสูงสุดที่ปลอดภัยคือไม่เกิน 6 เม็ดต่อวัน และไม่เกิน 10 เม็ดต่อสัปดาห์ เพื่อป้องกัน Ergotism และภาวะ Medication Overuse Headache (MOH)",
            "",
            "🔍 ข้ออื่นผิดเพราะอะไร:",
            "• ข้อ ก., ข., ค., ง. ล้วนเป็นข้อความที่ถูกต้องตามมาตรฐานการใช้ยา Ergotamine/Caffeine จึงไม่ใช่คำตอบของคำถามที่ถามหา 'ข้อใดไม่ถูกต้อง'",
            "",
            "📖 Guideline อ้างอิง:",
            "แนวทางการรักษาโรคปวดศีรษะไมเกรนสำหรับแพทย์ พ.ศ. 2563 (สมาคมประสาทวิทยาแห่งประเทศไทย) & American Headache Society (AHS) Guidelines on Acute Migraine Treatment",
            "",
            "📌 จุดจำก่อนสอบ:",
            "• Dosing Rules for Ergotamine: เริ่มต้น 1-2 เม็ด ซ้ำได้ 1 เม็ดทุก 30 นาที (Max 6 เม็ด/วัน, Max 10 เม็ด/สัปดาห์)",
            "• Ergotism Warning: ชา/ซีด/เย็นปลายมือปลายเท้า ต้องหยุดยาทันที",
            "• Contraindications: ห้ามใช้ในโรคหลอดเลือดหัวใจ (CAD), ความดันโลหิตสูงควบคุมไม่ได้, โรคหลอดเลือดสมอง, หญิงตั้งครรภ์/ให้นมบุตร (Uterine contraction) และห้ามใช้ร่วมกับ Strong CYP3A4 inhibitors (เช่น Clarithromycin, Ketoconazole, Ritonavir) เด็ดขาดเนื่องจากเสี่ยงต่อ Gangrene"
        ]
        ws.cell(r, 10).value = "\n".join(exp_lines)
        print(f"Updated Row {r} to Ans 5!")
        break

wb.save('PLE CC QUIZ.xlsx')
print("Saved successfully!")
