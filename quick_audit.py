# -*- coding: utf-8 -*-
import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx', read_only=True)
sheets = [
    '10. Psychiatric',
    '11. Pulmonary',
    '12. GynaecologicGenitourinary',
    '13. Eye disorder',
    '14. Oncologic',
    '15. Renal',
    '17. Herbal Medicine',
    '18. Clinical Nutrition'
]

for name in sheets:
    if name not in wb.sheetnames:
        print(f"Sheet {name} not found")
        continue
    ws = wb[name]
    template_count = 0
    asterisk_count = 0
    total_q = 0
    for r in range(3, ws.max_row + 1):
        q = ws.cell(r, 2).value
        if not q:
            continue
        total_q += 1
        exp = str(ws.cell(r, 10).value or '')
        if '*' in exp:
            asterisk_count += 1
        if 'ไม่ใช่คำตอบที่ถูกต้อง' in exp or 'ไม่ตรงกับบริบท' in exp:
            template_count += 1
    print(f"{name:30} | Questions: {total_q:4} | With Template: {template_count:4} | With Asterisk: {asterisk_count:4}")

wb.close()
