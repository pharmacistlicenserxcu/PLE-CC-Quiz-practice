import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx', read_only=True)

clinic_sheets = [
    '2. Cardiovascular', '3. Dermatologic', '4. Endocrine', '5. Gastrointestinal',
    '6. Hematologic', '7. Immunologic', '8. Infectious diseases', '9. Neurologic',
    '10. Psychiatric', '11. Pulmonary', '12. GynaecologicGenitourinary', '13. Eye disorder',
    '14. Oncologic', '15. Renal', '16. Others & Toxic', '17. Herbal Medicine', '18. Clinical Nutrition'
]

print(f"{'Sheet':<30} | {'Total Non-Empty':<16} | {'Asterisks (**)':<14} | {'Template Distractors'}")
print("-" * 85)

for s in clinic_sheets:
    if s not in wb.sheetnames:
        continue
    ws = wb[s]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows[2:] if any(x is not None for x in r)]
    
    asterisk_count = 0
    template_count = 0
    
    for r in non_empty:
        exp = str(r[9]) if len(r) > 9 and r[9] else ''
        if '**' in exp or '*' in exp:
            asterisk_count += 1
        if 'ไม่ใช่คำตอบที่ถูกต้อง' in exp or 'ไม่ตรงกับบริบท' in exp or 'ไม่ถูกต้องตามหลักการรักษา' in exp:
            template_count += 1
            
    print(f"{s:<30} | {len(non_empty):<16} | {asterisk_count:<14} | {template_count}")
