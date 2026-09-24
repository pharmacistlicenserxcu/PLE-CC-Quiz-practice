import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx')

sheets_to_update = [
    '🔍 รวมข้อสอบ & กรองข้อมูล', '1. Musculoskeleton', '2. Cardiovascular', '3. Dermatologic',
    '4. Endocrine', '5. Gastrointestinal', '6. Hematologic', '7. Immunologic',
    '8. Infectious diseases', '9. Neurologic', '10. Psychiatric', '11. Pulmonary',
    '12. GynaecologicGenitourinary', '13. Eye disorder', '14. Oncologic', '15. Renal',
    '6. Solid Dosage Forms', '7. Liquid & Semisolids', '13. Medicinal Chemistry',
    '1. Pharmacy Laws & Ethics', '2. Pharmacy Administration & Sy', '3. Research & Biostats',
    '18. Clinical Nutrition'
]

total_replaced = 0

for sname in sheets_to_update:
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    sheet_rep = 0
    for row in ws.iter_rows(min_row=3):
        # Col N (idx 14: หมายเหตุ), Col O (idx 15: ประเภทข้อสอบ), Col P (idx 16: ปี / เลขชุด)
        for col_idx in [14, 15, 16]:
            if len(row) >= col_idx:
                cell = row[col_idx - 1]
                if cell.value and 'ชุด 1' in str(cell.value):
                    orig = str(cell.value)
                    # Replace with 'Pharma Plus เล่มม่วง'
                    new_val = orig.replace('ชุด 1', 'Pharma Plus เล่มม่วง')
                    cell.value = new_val
                    sheet_rep += 1
    total_replaced += sheet_rep
    print(f"Sheet {sname:<30}: replaced {sheet_rep} occurrences")

print("-" * 60)
print(f"Total cells updated: {total_replaced}")
wb.save('PLE CC QUIZ.xlsx')
print("Saved PLE CC QUIZ.xlsx successfully.")
