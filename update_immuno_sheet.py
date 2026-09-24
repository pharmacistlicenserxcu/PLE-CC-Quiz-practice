import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import json
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Load master pristine immunologic questions
with open('immuno_master_pristine.json', 'r', encoding='utf-8') as f:
    items = json.load(f)

print(f"Loaded {len(items)} pristine immunologic items.")

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx')
sheet_name = '7. Immunologic'
if sheet_name not in wb.sheetnames:
    print(f"Error: {sheet_name} not found in workbook!")
    sys.exit(1)

ws = wb[sheet_name]

# Ensure banner in Row 1
ws['A1'] = '=HYPERLINK("#gid=0", "🏠 กลับสู่หน้าแรก (Go to Home Page)")'
ws['A1'].font = Font(name='Bai Jamjuree', size=11, bold=True, color='0D47A1')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws['A1'].fill = PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid')
ws.row_dimensions[1].height = 35

# Keep headers in Row 2
# Clear all rows from Row 3 to max_row
max_r = ws.max_row
print(f"Current max_row in sheet: {max_r}")

for row in ws.iter_rows(min_row=3, max_row=max_r):
    for cell in row:
        cell.value = None

thin_border = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)

font_regular = Font(name='Bai Jamjuree', size=10)
align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
align_center = Alignment(horizontal='center', vertical='top')

for idx, it in enumerate(items, 1):
    row_num = idx + 2
    
    # Col A: ข้อที่
    ws.cell(row=row_num, column=1, value=idx).alignment = align_center
    ws.cell(row=row_num, column=1).font = font_regular
    
    # Col B: คำถาม
    ws.cell(row=row_num, column=2, value=it.get('q_text', '')).alignment = align_left
    ws.cell(row=row_num, column=2).font = font_regular
    
    # Col C: รูปถาม
    ws.cell(row=row_num, column=3, value=it.get('q_img', '')).alignment = align_center
    ws.cell(row=row_num, column=3).font = font_regular
    
    # Col D-H: ตัวเลือก 1-5
    for c_idx in range(1, 6):
        col_no = c_idx + 3
        c_val = it.get(f'c{c_idx}', '')
        ws.cell(row=row_num, column=col_no, value=c_val).alignment = align_left
        ws.cell(row=row_num, column=col_no).font = font_regular
        
    # Col I: เฉลย (1-5)
    ws.cell(row=row_num, column=9, value=it.get('ans', '')).alignment = align_center
    ws.cell(row=row_num, column=9).font = font_regular
    
    # Col J: คำอธิบายเฉลย
    ws.cell(row=row_num, column=10, value=it.get('exp', '')).alignment = align_left
    ws.cell(row=row_num, column=10).font = font_regular
    
    # Col K: รูปเฉลย
    ws.cell(row=row_num, column=11, value=it.get('exp_img', '')).alignment = align_center
    ws.cell(row=row_num, column=11).font = font_regular
    
    # Col L: Filter หมวด/Subtopic
    ws.cell(row=row_num, column=12, value=it.get('subtopic', '')).alignment = align_left
    ws.cell(row=row_num, column=12).font = font_regular
    
    # Col M: Product / Clinic
    ws.cell(row=row_num, column=13, value='Clinic').alignment = align_center
    ws.cell(row=row_num, column=13).font = font_regular
    
    # Col N: หมายเหตุ
    ws.cell(row=row_num, column=14, value=it.get('note', '')).alignment = align_left
    ws.cell(row=row_num, column=14).font = font_regular
    
    # Col O: ประเภทข้อสอบ
    ws.cell(row=row_num, column=15, value=it.get('q_type', '')).alignment = align_center
    ws.cell(row=row_num, column=15).font = font_regular
    
    # Col P: ปี / เลขชุด
    ws.cell(row=row_num, column=16, value=it.get('exam_year', '')).alignment = align_center
    ws.cell(row=row_num, column=16).font = font_regular

print(f"Saving changes to PLE CC QUIZ.xlsx...")
wb.save('PLE CC QUIZ.xlsx')
print(f"Successfully wrote {len(items)} rows to 7. Immunologic.")
