import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx', read_only=True)
print("Sheet 11. Herbal Products headers:")
ws_herb = wb['11. Herbal Products']
for r in list(ws_herb.iter_rows(values_only=True))[:3]:
    print(" ", r[:6])

print("\nSheet 12. Food Products & QA headers:")
ws_food = wb['12. Food Products & QA']
for r in list(ws_food.iter_rows(values_only=True))[:3]:
    print(" ", r[:6])
