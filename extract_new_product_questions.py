import sys
import io
import re
import json
import openpyxl

# Do not wrap sys.stdout again
from classify_engine import classify_product_question

src = r'C:\Users\thana\Downloads\_CC1 Product Contents (1).md'
with open(src, 'r', encoding='utf-8') as f:
    text = f.read()

# Load existing questions from PLE CC QUIZ.xlsx to prevent duplication
wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx', read_only=True)
existing_questions = set()
for sname in wb.sheetnames:
    ws = wb[sname]
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r and len(r) > 1 and r[1]:
            q_clean = str(r[1]).strip()[:50]
            existing_questions.add(q_clean)

sys.stdout.write(f"Loaded {len(existing_questions)} existing question signatures from workbook.\n")

split_idx = text.find('# แบ่งงาน')
if split_idx == -1:
    split_idx = len(text)
exam_text = text[:split_idx]

year_spans = [
    ('2567', 47, 25508),
    ('2566', 25508, 43414),
    ('2565', 43414, 75498),
    ('2564', 75498, 189448),
    ('2563', 189448, 230633),
    ('2562', 230633, 299916),
    ('2561', 299916, 377373),
    ('2560', 377373, 379117),
    ('2559', 379117, 390690),
    ('2558', 390690, 447909)
]

def clean_text(s):
    if not s:
        return ""
    s = str(s).replace('**', '').replace('*', '').replace('_x001E_', '')
    s = s.replace('\\-', '-').replace('\\.', '.').replace('&nbsp;', ' ')
    return s.strip()

parsed_new_questions = []

for yr, start_pos, end_pos in year_spans:
    chunk = exam_text[start_pos:end_pos]
    
    if 'ส่วนที่ 2: เฉลย' in chunk or 'ส่วนที่ 2: Answer' in chunk or 'Answer Key & Explanations' in chunk:
        parts = re.split(r'ส่วนที่ 2:\s*(?:เฉลยข้อสอบ|Answer|Answer Key & Explanations)', chunk)
        q_chunk = parts[0]
        ans_chunk = parts[1] if len(parts) > 1 else ""
    else:
        q_chunk = chunk
        ans_chunk = ""

    def replace_img(match):
        img_id = match.group(1)
        return f"[แนบรูป: ข้อสอบ Product ปี {yr} รหัสรูปภาพ {img_id}]"

    q_chunk_clean = re.sub(r'!\[.*?\]\[(.*?)\]', replace_img, q_chunk)
    ans_chunk_clean = re.sub(r'!\[.*?\]\[(.*?)\]', replace_img, ans_chunk)

    q_matches = list(re.finditer(r'(?:### \*\*ข้อ\s*(\d+)\*\*|\bข้อที่?\s*(\d+)|\b(\d+)\\\.\s+)(.*?)(?=(?:### \*\*ข้อ|\bข้อที่?\s*\d+|\b\d+\\\.\s+|$))', q_chunk_clean, re.DOTALL))
    
    sys.stdout.write(f"Year {yr}: found {len(q_matches)} raw question candidates.\n")
    
    for qm in q_matches:
        q_num = qm.group(1) or qm.group(2) or qm.group(3)
        raw_content = qm.group(4).strip()
        lines = [l.strip() for l in raw_content.splitlines() if l.strip()]
        if not lines:
            continue
            
        q_lines = []
        choices = {}
        curr_choice = None
        has_choices = False
        
        for l in lines:
            c_m = re.match(r'^(?:\*\*)?([ก-จa-eA-E1-5])[\.\)]\s*(.*)', l)
            if c_m and ('ตอบ:' not in l and 'เฉลย:' not in l):
                has_choices = True
                curr_choice = c_m.group(1).upper()
                t_map = {'A': 'ก', 'B': 'ข', 'C': 'ค', 'D': 'ง', 'E': 'จ', '1': 'ก', '2': 'ข', '3': 'ค', '4': 'ง', '5': 'จ'}
                curr_choice = t_map.get(curr_choice, curr_choice)
                choices[curr_choice] = clean_text(c_m.group(2))
            elif curr_choice and has_choices and not l.startswith('###') and not l.startswith('ตอบ'):
                choices[curr_choice] += " " + clean_text(l)
            else:
                if not has_choices:
                    q_lines.append(l)
                    
        q_full = clean_text(" ".join(q_lines))
        if q_full.startswith('โจทย์:'):
            q_full = q_full[6:].strip()
            
        if len(q_full) < 15 or q_full[:50] in existing_questions:
            continue
            
        ans_num = 1
        ans_match = re.search(r'ตอบ:?\s*([ก-จa-eA-E1-5])', raw_content)
        if ans_match:
            a_key = ans_match.group(1).upper()
            t_map = {'A': 1, 'B': 2, 'C': 3, 'D': 4, 'E': 5, 'ก': 1, 'ข': 2, 'ค': 3, 'ง': 4, 'จ': 5, '1': 1, '2': 2, '3': 3, '4': 4, '5': 5}
            ans_num = t_map.get(a_key, 1)

        q_img_tag = None
        img_match = re.search(r'\[แนบรูป: [^\]]+\]', raw_content)
        if img_match:
            q_img_tag = img_match.group(0)

        c_list = [choices.get('ก', ''), choices.get('ข', ''), choices.get('ค', ''), choices.get('ง', ''), choices.get('จ', '')]
        target_tab, subtopic_name = classify_product_question(q_full, c_list, raw_content)
        
        parsed_new_questions.append({
            'year': yr,
            'q_num': q_num,
            'q_text': q_full,
            'q_img': q_img_tag,
            'c1': choices.get('ก', ''),
            'c2': choices.get('ข', ''),
            'c3': choices.get('ค', ''),
            'c4': choices.get('ง', ''),
            'c5': choices.get('จ', ''),
            'ans': ans_num,
            'raw_content': raw_content,
            'target_tab': target_tab,
            'subtopic': subtopic_name
        })

sys.stdout.write(f"\nTotal new extracted questions across 10 years: {len(parsed_new_questions)}\n")

tab_counts = {}
for q in parsed_new_questions:
    tab_counts[q['target_tab']] = tab_counts.get(q['target_tab'], 0) + 1

sys.stdout.write("\nBreakdown of new questions to ingest by Target Tab:\n")
for t, cnt in sorted(tab_counts.items(), key=lambda x: -x[1]):
    sys.stdout.write(f"  - {t:<30}: {cnt} questions\n")

with open('parsed_new_product_questions.json', 'w', encoding='utf-8') as f:
    json.dump(parsed_new_questions, f, ensure_ascii=False, indent=2)

sys.stdout.write("Saved to parsed_new_product_questions.json\n")
