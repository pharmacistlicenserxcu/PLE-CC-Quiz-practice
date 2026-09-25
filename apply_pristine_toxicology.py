#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rebuild Sheet 16 as Pristine Toxicology & Poisoning Exam Bank
Sheet Name: '16. Toxicology & Poisoning' (formerly '16. Others & Toxic')
Enforces:
1. Exact GEMINI.md Banner at Row 1 (A1:F1 merged, Bai Jamjuree 11pt Bold, 0D47A1 on E3F2FD)
2. Row 2 Standard Headers
3. 40 High-Yield Pristine Clinical Toxicology Exam Questions
4. Strict zero asterisks, zero templates, exact answer letter/num matching
5. Update 'สารบัญ' (Index) sheet to reflect the new pristine sheet title
"""

import sys, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

EXCEL_FILE = "PLE CC QUIZ.xlsx"
OLD_SHEET = "16. Others & Toxic"
NEW_SHEET = "16. Toxicology & Poisoning"

print(f"Loading {EXCEL_FILE}...")
wb = openpyxl.load_workbook(EXCEL_FILE)

# 1. Rename or get sheet
if OLD_SHEET in wb.sheetnames:
    ws = wb[OLD_SHEET]
    ws.title = NEW_SHEET
elif NEW_SHEET in wb.sheetnames:
    ws = wb[NEW_SHEET]
else:
    ws = wb.create_sheet(NEW_SHEET)

print(f"Working on sheet: {ws.title}")

# 2. Clear all rows
for r in range(1, ws.max_row + 5):
    for c in range(1, 16):
        try:
            ws.cell(r, c).value = None
        except:
            pass

# Clear merges
for merge in list(ws.merged_cells.ranges):
    try:
        ws.unmerge_cells(str(merge))
    except:
        pass

# 3. Apply Banner at Row 1
BANNER_FONT = Font(name="Bai Jamjuree", size=11, bold=True, color="0D47A1")
BANNER_FILL = PatternFill(fill_type="solid", fgColor="E3F2FD")
BANNER_ALIGN = Alignment(horizontal="center", vertical="center")

ws.row_dimensions[1].height = 35
ws["A1"] = '=HYPERLINK("#gid=0", "\U0001f3e0 กลับสู่หน้าแรก (Go to Home Page)")'
ws["A1"].font = BANNER_FONT
ws["A1"].fill = BANNER_FILL
ws["A1"].alignment = BANNER_ALIGN
ws.merge_cells("A1:F1")

# 4. Write Row 2 Headers
headers = [
    'ข้อที่', 'คำถาม', 'รูปถาม', 'ตัวเลือก 1', 'ตัวเลือก 2', 'ตัวเลือก 3',
    'ตัวเลือก 4', 'ตัวเลือก 5', 'เฉลย (ตัวเลข 1-5)', 'คำอธิบายเฉลย',
    'รูปเฉลย', 'Filter หมวด/Subtopic', 'Product / Clinic', 'หมายเหตุ'
]

HEADER_FONT = Font(name="Bai Jamjuree", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1565C0") # Deep Blue
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws.row_dimensions[2].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(2, col_idx, h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = HEADER_ALIGN

# 5. Load 40 Pristine Questions
data = json.load(open('pristine_toxicology_40.json', encoding='utf-8'))

DATA_FONT = Font(name="Bai Jamjuree", size=10)
DATA_ALIGN_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
DATA_ALIGN_CENTER = Alignment(horizontal="center", vertical="top")

for idx, it in enumerate(data, 1):
    row_num = idx + 2
    ws.row_dimensions[row_num].height = 65
    
    ws.cell(row_num, 1, idx).alignment = DATA_ALIGN_CENTER
    ws.cell(row_num, 2, it['q']).alignment = DATA_ALIGN_LEFT
    ws.cell(row_num, 3, "")
    ws.cell(row_num, 4, it['c1']).alignment = DATA_ALIGN_LEFT
    ws.cell(row_num, 5, it['c2']).alignment = DATA_ALIGN_LEFT
    ws.cell(row_num, 6, it['c3']).alignment = DATA_ALIGN_LEFT
    ws.cell(row_num, 7, it['c4']).alignment = DATA_ALIGN_LEFT
    ws.cell(row_num, 8, it.get('c5', '')).alignment = DATA_ALIGN_LEFT
    ws.cell(row_num, 9, it['ans']).alignment = DATA_ALIGN_CENTER
    ws.cell(row_num, 10, it['exp']).alignment = DATA_ALIGN_LEFT
    ws.cell(row_num, 11, "")
    ws.cell(row_num, 12, it.get('subtopic', 'Toxicology')).alignment = DATA_ALIGN_CENTER
    ws.cell(row_num, 13, it.get('cat', 'Clinic')).alignment = DATA_ALIGN_CENTER
    ws.cell(row_num, 14, it.get('note', '')).alignment = DATA_ALIGN_LEFT
    
    for c in range(1, 15):
        ws.cell(row_num, c).font = DATA_FONT

# 6. Update สารบัญ (Index Sheet) if exists
if 'สารบัญ' in wb.sheetnames:
    ws_idx = wb['สารบัญ']
    for r in range(1, ws_idx.max_row + 1):
        for c in range(1, ws_idx.max_column + 1):
            val = str(ws_idx.cell(r, c).value or '')
            if '16. Others & Toxic' in val or 'Others & Toxic' in val:
                new_val = val.replace('16. Others & Toxic', '16. Toxicology & Poisoning').replace('Others & Toxic', 'Toxicology & Poisoning')
                ws_idx.cell(r, c).value = new_val
                print(f"Updated index at R{r}C{c}: {new_val}")

# Set column widths
col_widths = {
    'A': 8, 'B': 45, 'C': 8, 'D': 30, 'E': 30, 'F': 30, 'G': 30, 'H': 30,
    'I': 14, 'J': 60, 'K': 8, 'L': 25, 'M': 15, 'N': 30
}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

print(f"Saving workbook {EXCEL_FILE}...")
wb.save(EXCEL_FILE)
print("Workbook saved successfully!")

# Final Verification
wb_check = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
ws_check = wb_check[NEW_SHEET]
print(f"\nAUDIT FOR [{NEW_SHEET}]:")
print(f"Max Row: {ws_check.max_row} (Questions: {ws_check.max_row - 2})")
print(f"Banner A1: {ws_check['A1'].value}")

num_to_let = {1: 'ก', 2: 'ข', 3: 'ค', 4: 'ง', 5: 'จ'}
mismatch = 0
asterisks = 0
templates = 0

for r in range(3, ws_check.max_row + 1):
    ans = int(ws_check.cell(r, 9).value)
    let = num_to_let[ans]
    exp = str(ws_check.cell(r, 10).value or '')
    first_line = exp.split('\n')[0]
    if f"ข้อ {let}." not in first_line:
        mismatch += 1
    if "*" in exp:
        asterisks += 1
    if "ไม่ใช่คำตอบที่ถูกต้องสำหรับภาวะ" in exp or "ไม่ตรงกับบริบท" in exp:
        templates += 1

print(f"Answer Mismatches: {mismatch}")
print(f"Asterisks Found: {asterisks}")
print(f"Template Explanations: {templates}")
print("STATUS: " + ("PASS ✅" if mismatch == 0 and asterisks == 0 and templates == 0 else "FAIL ❌"))
