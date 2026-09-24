import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx', read_only=True)
found_sheets = {}

for sname in wb.sheetnames:
    ws = wb[sname]
    sheet_count = 0
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r and len(r) > 13:
            for c_idx in range(13, min(len(r), 16)):
                val = str(r[c_idx]) if r[c_idx] else ''
                if 'ชุด 1' in val:
                    sheet_count += 1
                    break
    if sheet_count > 0:
        found_sheets[sname] = sheet_count

print("Total occurrences of 'ชุด 1':")
total_rows = 0
for s, c in found_sheets.items():
    print(f"  - {s:<30}: {c} rows")
    total_rows += c

print("-" * 50)
print(f"Total across all sheets: {total_rows} rows")
