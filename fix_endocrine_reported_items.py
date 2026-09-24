import openpyxl
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx')
ws = wb['4. Endocrine']

print('--- Fixing Row 7 (Ref 5) ---')
# Question: ข้อใด ไม่ใช่ ภาวะแทรกซ้อนระยะยาวของโรคเบาหวาน
# Choices: ก. Retinopathy, ข. Nephropathy, ค. Neuropathy, ง. Glaucoma, จ. Atherosclerosis
ws.cell(7, 9, 4)
exp_7 = """✅ คำตอบที่ถูกต้อง: ข้อ ง. Glaucoma

💡 Background:
ภาวะแทรกซ้อนระยะยาวของโรคเบาหวาน (Chronic complications of DM) แบ่งออกเป็น 2 กลุ่มหลัก คือ Microvascular complications (Retinopathy, Nephropathy, Neuropathy) และ Macrovascular complications (Coronary artery disease, Peripheral artery disease, Cerebrovascular disease/Atherosclerosis)

🎯 ทำไมข้อนี้ถึงถูก:
Glaucoma (โรคต้อหิน) ไม่ใช่ภาวะแทรกซ้อนระยะยาวแบบปฐมภูมิหรือ classic vascular complication ของโรคเบาหวาน แม้ว่าผู้ป่วยเบาหวานอาจมีความเสี่ยงต่อ Neovascular glaucoma ในระยะท้ายของ Proliferative diabetic retinopathy ได้ แต่ภาวะแทรกซ้อนทางตาหลักของเบาหวานคือ Diabetic Retinopathy และ Cataract ไม่ใช่ Glaucoma

🔍 ข้ออื่นผิดเพราะอะไร:
• ข้อ ก. (Retinopathy): เป็น Microvascular complication ที่พบบ่อย เกิดจาก microaneurysms, vascular permeability และ neovascularization ที่จอประสาทตา
• ข้อ ข. (Nephropathy): เป็น Microvascular complication เด่นชัด เกิดจาก glomerular sclerosis, hyperfiltration และ persistent albuminuria
• ข้อ ค. (Neuropathy): เป็น Microvascular complication ที่ส่งผลต่อเส้นประสาทส่วนปลาย (Distal symmetric polyneuropathy) และระบบประสาทอัตโนมัติ (Autonomic neuropathy)
• ข้อ จ. (Atherosclerosis): เป็น Macrovascular complication หลักที่เกิดจากหลอดเลือดแดงแข็งตัว เพิ่มความเสี่ยงต่อ Myocardial infarction, Stroke และ Peripheral arterial disease

📖 Guideline อ้างอิง:
American Diabetes Association (ADA) Standards of Care in Diabetes 2024 & แนวทางเวชปฏิบัติสำหรับโรคเบาหวาน พ.ศ. 2566 (สมาคมโรคเบาหวานแห่งประเทศไทย)

📌 จุดจำก่อนสอบ:
• Microvascular Triad: Retinopathy, Nephropathy, Neuropathy สัมพันธ์โดยตรงกับระดับ glycemic control (HbA1c)
• Macrovascular Complications: สัมพันธ์กับปัจจัยเสี่ยงร่วม เช่น ความดันโลหิตสูง ไขมันในเลือดผิดปกติ และการสูบบุหรี่
• Eye Examination: T2DM ควรตรวจ dilated eye exam ทันทีที่ได้รับการวินิจฉัย และตรวจซ้ำอย่างน้อยปีละ 1 ครั้ง"""
ws.cell(7, 10, exp_7)

print('--- Fixing Row 50 (Ref 48) ---')
ws.cell(50, 8, 'Metformin')
ws.cell(50, 9, 5)
exp_50 = """✅ คำตอบที่ถูกต้อง: ข้อ จ. Metformin

💡 Background:
การลดภาวะแทรกซ้อนทางหลอดเลือดขนาดใหญ่ (Macrovascular complications เช่น Myocardial infarction, Stroke) ในผู้ป่วยเบาหวานชนิดที่ 2 ต้องอาศัยการคุมระดับน้ำตาลควบคู่กับการจัดการปัจจัยเสี่ยงหลอดเลือดหัวใจ

🎯 ทำไมข้อนี้ถึงถูก:
Metformin จากการศึกษาประวัติศาสตร์ระยะยาว UKPDS (United Kingdom Prospective Diabetes Study) และ UKPDS 10-year post-trial follow-up พบว่า Metformin ในผู้ป่วย T2DM ที่มีภาวะน้ำหนักเกิน/อ้วน สามารถลดอุบัติการณ์ของ Myocardial infarction และ All-cause mortality ได้อย่างมีนัยสำคัญทางสถิติเมื่อเทียบกับ conventional therapy ถือเป็นยารับประทานดั้งเดิมที่มีหลักฐาน cardiovascular benefit ชัดเจนที่สุด

🔍 ข้ออื่นผิดเพราะอะไร:
• ข้อ ก. (Acarbose และ metformin): แม้จะมี Metformin แต่ Acarbose ซึ่งเป็น Alpha-glucosidase inhibitor มีหลักฐานการลด macrovascular outcomes ไม่เด่นชัดเท่า Metformin เดี่ยวๆ หรือยากลุ่มใหม่ SGLT2i/GLP-1RA
• ข้อ ข. (Glipizide และ pioglitazone): Glipizide ซึ่งเป็น Sulfonylurea ไม่มีผลลด macrovascular death และเสี่ยงต่อ hypoglycemia ส่วน Pioglitazone เพิ่มความเสี่ยงต่อภาวะน้ำคั่งและ Heart Failure กำเริบ
• ข้อ ค. (SU และ metformin): Sulfonylurea ไม่ได้ลด macrovascular events และการเกิด hypoglycemia สัมพันธ์กับ cardiovascular mortality ที่สูงขึ้น
• ข้อ ง. (Repaglinide และ SU): ยากลุ่มกระตุ้นการหลั่งอินซูลินทั้งสองกลุ่มไม่มีหลักฐานทางคลินิกที่ยืนยันการลดภาวะแทรกซ้อนทางหลอดเลือดขนาดใหญ่

📖 Guideline อ้างอิง:
American Diabetes Association (ADA) Standards of Care in Diabetes 2024 & UKPDS 34 Study (Lancet 1998)

📌 จุดจำก่อนสอบ:
• UKPDS Legacy Effect: Metformin แสดงผลลด MI และ total mortality ต่อเนื่องยาวนานแม้หลังจบการศึกษา
• Modern Cardioprotection: ในผู้ป่วยที่มี ASCVD หรือ high CV risk ปัจจุบันแนะนำเพิ่ม SGLT2i (Empagliflozin, Dapagliflozin) หรือ GLP-1 RA (Semaglutide, Liraglutide, Dulaglutide) ที่มี proven CV benefit ร่วมด้วย"""
ws.cell(50, 10, exp_50)

print('--- Fixing Row 85 (Ref 83) ---')
ws.cell(85, 8, 'หากมีผื่นขึ้นแม้เพียงเล็กน้อย ให้หยุดยาทันที')
ws.cell(85, 9, 4)
exp_85 = """✅ คำตอบที่ถูกต้อง: ข้อ ง. หากมีอาการ ไข้ เจ็บคอ มีแผลในปาก ควรหยุดยาแล้วมาพบแพทย์ทันที

💡 Background:
ยาต้านไทรอยด์กลุ่ม Thionamides (ได้แก่ Methimazole/MMI และ Propylthiouracil/PTU) เป็นการรักษาหลักสำหรับ Graves' disease แต่มีอาการไม่พึงประสงค์ที่อันตรายถึงชีวิตคือ Agranulocytosis (เม็ดเลือดขาวชนิดนิวโทรฟิลต่ำรุนแรง)

🎯 ทำไมข้อนี้ถึงถูก:
หากมีอาการ ไข้ เจ็บคอ มีแผลในปาก ควรหยุดยาแล้วมาพบแพทย์ทันที ตาม 2018 European Thyroid Association (ETA) Guideline และ 2016 American Thyroid Association (ATA) Guideline กำหนดให้เภสัชกรและแพทย์ต้องเตือนผู้ป่วยทุกคนที่เริ่มยา MMI/PTU ว่า หากมีอาการติดเชื้อเฉียบพลัน เช่น ไข้ เจ็บคอ (Sore throat) มีแผลในปาก (Mouth ulcers) ต้อง หยุดยาทันที และรีบมาโรงพยาบาลเพื่อตรวจความสมบูรณ์ของเม็ดเลือด (CBC with differential) ทันทีเพื่อตรวจคัดกรองภาวะ Agranulocytosis (ANC < 500 cells/mm³)

🔍 ข้ออื่นผิดเพราะอะไร:
• ข้อ ก. (ให้กินยาจนครบคอร์ส 6-8 เดือน ห้ามหยุดยาเอง ก่อนปรึกษาแพทย์): ระยะเวลาการรักษา Graves' disease ด้วยยาต้านไทรอยด์ตาม guideline สากลคือ 12-18 เดือน (ไม่ใช่ 6-8 เดือน) เพื่อลดอัตราการกลับเป็นซ้ำ (relapse)
• ข้อ ข. (ยาอาจทำให้ง่วงซึม หลีกเลี่ยงการขับขี่ยานพาหนะหรือทำงานเครื่องจักร): MMI และ PTU ไม่ได้ออกฤทธิ์กดประสาทส่วนกลางและไม่ทำให้ง่วงซึม
• ข้อ ค. (หลีกเลี่ยงการออกแดดหรือให้ทาครีมกันแดด เมื่อจำเป็น): ไม่ใช่อาการไม่พึงประสงค์หลักหรือคำเตือนสำคัญของยาต้านไทรอยด์
• ข้อ จ. (หากมีผื่นขึ้นแม้เพียงเล็กน้อย ให้หยุดยาทันที): หากมีผื่นแพ้เล็กน้อย (Minor allergic skin rash / urticaria) สามารถให้ยาแก้แพ้ (Antihistamines) บรรเทาอาการร่วมได้โดยไม่จำเป็นต้องหยุดยาต้านไทรอยด์ทันที ยกเว้นกรณีผื่นแพ้รุนแรง เช่น SJS/TEN หรือมีอาการ systemic involvement

📖 Guideline อ้างอิง:
2018 European Thyroid Association (ETA) Guideline for the Management of Graves’ Hyperthyroidism & 2016 American Thyroid Association (ATA) Guidelines for Diagnosis and Management of Hyperthyroidism

📌 จุดจำก่อนสอบ:
• Critical Warning: ไข้ + เจ็บคอ = หยุดยาทันที + ตรวจ CBC ฉุกเฉินเพื่อเช็ก Agranulocytosis
• Duration of Therapy: รักษาด้วย MMI ต่อเนื่อง 12-18 เดือน แล้วจึงประเมินการหยุดยาเมื่อ TRAb ให้ผลลบและ TSH ปกติ
• Major Toxicities: MMI เสี่ยงต่อ Cholestatic jaundice, PTU เสี่ยงต่อ Fulminant hepatotoxicity (Black box warning)"""
ws.cell(85, 10, exp_85)

print('--- Fixing Row 86 (Ref 84) ---')
q_86 = """[สถานการณ์: ผู้ป่วยหญิง 55 ปี โรคประจำตัวเบาหวาน ความดันโลหิตสูง ได้รับ Metformin 500 mg 1x2 pc, Glipizide 5 mg 1x1 pc, ผลตรวจ HbA1c ล่าสุด 8.1%, eGFR: 48-49 mL/min/1.73 m², UACR: 282 mg/g (CKD Stage 3aA2)]
ยาเบาหวานที่เหมาะสมในผู้ป่วยรายนี้ ตามแนวทางเวชปฏิบัติ KDIGO 2022 คือข้อใด?"""
ws.cell(86, 2, q_86)
ws.cell(86, 8, 'Metformin + sitagliptin')
ws.cell(86, 9, 1)
exp_86 = """✅ คำตอบที่ถูกต้อง: ข้อ ก. Metformin + dapagliflozin

💡 Background:
แนวทางเวชปฏิบัติ KDIGO 2022 Clinical Practice Guideline for Diabetes Management in Chronic Kidney Disease แนะนำการดูแลรักษาผู้ป่วยเบาหวานร่วมกับโรคไตเรื้อรัง (CKD) โดยเน้นยาที่มีการชะลอการเสื่อมของไต (Kidney protection) และลดความเสี่ยงโรคหลอดเลือดหัวใจ

🎯 ทำไมข้อนี้ถึงถูก:
Metformin + dapagliflozin ตาม KDIGO 2022 แนะนำให้ First-line pharmacotherapy สำหรับผู้ป่วย T2DM ร่วมกับ CKD ที่มี eGFR ≥ 30 mL/min/1.73 m² ประกอบด้วย:
1) Metformin (ใช้ได้อย่างปลอดภัยเมื่อ eGFR ≥ 45 mL/min/1.73 m² แต่ควรปรับลดขนาดยาเหลือไม่เกิน 1,000 mg/วัน หาก eGFR 30-44)
2) SGLT2 inhibitor (Dapagliflozin หรือ Empagliflozin) แนะนำให้เริ่มในผู้ป่วยที่มี eGFR ≥ 20 mL/min/1.73 m² และมี Albuminuria (UACR > 30 mg/g) เพื่อชะลอการดำเนินโรคมุ่งสู่ไตวายเรื้อรังระยะสุดท้ายและลด cardiovascular events
ผู้ป่วยรายนี้มี eGFR 48-49 และ UACR 282 mg/g การให้ Metformin ร่วมกับ SGLT2i (Dapagliflozin) จึงตรงตามเป้าหมายสูงสุดของ KDIGO 2022

🔍 ข้ออื่นผิดเพราะอะไร:
• ข้อ ข. (Glipizide + insulin): Glipizide แม้จะปลอดภัยในไตเพราะขับออกทางตับ แต่ไม่มีหลักฐานชะลอการเสื่อมของไต และทั้งสองตัวเพิ่มความเสี่ยงต่อการเกิด Hypoglycemia อย่างมีนัยสำคัญ
• ข้อ ค. (Insulin): เป็นทางเลือกเสริมเมื่อยาอื่นไม่สามารถคุมระดับน้ำตาลได้ตามเป้าหมาย หรือ eGFR < 30 mL/min แต่ไม่มีฤทธิ์จำเพาะในการชะลอการเสื่อมของไต (No direct nephroprotection)
• ข้อ ง. (Linagliptin): DPP-4 inhibitor เช่น Linagliptin มีความปลอดภัยสูง ไม่ต้องปรับขนาดยาตามการทำงานของไต แต่ตาม KDIGO 2022 จัดเป็น Second-line (Non-first-line) รองจาก Metformin + SGLT2i
• ข้อ จ. (Metformin + sitagliptin): Sitagliptin ต้องปรับลดยาตาม eGFR และไม่มีข้อมูลในการชะลอโรคไตเรื้อรังเมื่อเทียบกับ SGLT2 inhibitors

📖 Guideline อ้างอิง:
KDIGO 2022 Clinical Practice Guideline for Diabetes Management in Chronic Kidney Disease & ADA Standards of Care in Diabetes 2024

📌 จุดจำก่อนสอบ:
• First-line KDIGO 2022: Lifestyle modification + Metformin + SGLT2 inhibitor (+ Statin + ACEI/ARB)
• SGLT2i Renal Threshold: สามารถเริ่มใช้ได้จนถึง eGFR ≥ 20 mL/min/1.73 m² และให้ยาต่อไปได้จนกว่าจะเริ่มฟอกไต
• Metformin eGFR Cutoffs: ห้ามเริ่มถ้า eGFR < 45, ลดขนาดยาสูงสุดเหลือ 1,000 mg/วัน ถ้า eGFR 30-44, และหยุดยาเด็ดขาดเมื่อ eGFR < 30 mL/min/1.73 m²"""
ws.cell(86, 10, exp_86)

print('--- Fixing Row 87 (Ref 85) ---')
q_87 = """[สถานการณ์: ผู้ป่วยตั้งครรภ์ 5 สัปดาห์ มีอาการกินจุแต่น้ำหนักลด เหงื่อออกมากผิดปกติ เหนื่อยง่าย ใจสั่น สงสัยภาวะ Hyperthyroidism จาก Graves' disease]
ผู้ป่วยรายนี้ ผลการตรวจฮอร์โมนต่อมไทรอยด์จากห้องปฏิบัติการควรมีลักษณะสอดคล้องกับข้อใด?"""
ws.cell(87, 2, q_87)
ws.cell(87, 9, 3)
exp_87 = """✅ คำตอบที่ถูกต้อง: ข้อ ค. TSH ต่ำ, FT3 สูง, FT4 สูง

💡 Background:
Graves' disease เป็นสาเหตุที่พบบ่อยที่สุดของภาวะต่อมไทรอยด์ทำงานเกิน (Primary hyperthyroidism) เกิดจากร่างกายสร้าง Thyroid-stimulating immunoglobulins (TSI / TRAb) ไปกระตุ้น TSH receptor บนต่อมไทรอยด์อย่างต่อเนื่อง

🎯 ทำไมข้อนี้ถึงถูก:
TSH ต่ำ, FT3 สูง, FT4 สูง ในภาวะ Primary hyperthyroidism จาก Graves' disease ต่อมไทรอยด์จะสังเคราะห์และหลั่งฮอร์โมน Free T3 (FT3) และ Free T4 (FT4) ออกมาในกระแสเลือดสูงกว่าปกติ ฮอร์โมนไทรอยด์ที่สูงนี้จะไปออกฤทธิ์ยับยั้งย้อนกลับแบบลบ (Negative feedback inhibition) ที่ต่อมใต้สมองส่วนหน้า (Anterior pituitary gland) ส่งผลให้ระดับ Thyroid-stimulating hormone (TSH) ลดต่ำลงจนวัดค่าแทบไม่ได้ (Suppressed TSH < 0.01 mIU/L)

🔍 ข้ออื่นผิดเพราะอะไร:
• ข้อ ก. (TSH สูง, FT3 ปกติ, FT4 ปกติ): เป็นลักษณะของ Subclinical hypothyroidism
• ข้อ ข. (TSH ปกติ, FT3 สูง, FT4 สูง): อาจพบได้ในภาวะ Artifact / Assay interference, TSH-secreting pituitary adenoma หรือ Thyroid hormone resistance syndrome ไม่ใช่ลักษณะของ Graves' disease
• ข้อ ง. (TSH ต่ำ, FT3 ต่ำ, FT4 ต่ำ): เป็นลักษณะของ Secondary / Central hypothyroidism (ความผิดปกติที่ระดับ Pituitary หรือ Hypothalamus)
• ข้อ จ. (TSH สูง, FT3 สูง, FT4 สูง): เป็นลักษณะของ Secondary hyperthyroidism (เช่น TSH-producing pituitary tumor) ซึ่งพบได้น้อยมาก ต่างจาก Graves' disease ซึ่งเป็น Primary disorder

📖 Guideline อ้างอิง:
2016 American Thyroid Association (ATA) Guidelines for Diagnosis and Management of Hyperthyroidism and Other Causes of Thyrotoxicosis

📌 จุดจำก่อนสอบ:
• Primary Hyperthyroidism Hallmark: TSH ต่ำ (Suppressed) ร่วมกับ FT4 และ/หรือ FT3 สูง
• Primary Hypothyroidism Hallmark: TSH สูง (Elevated) ร่วมกับ FT4 ต่ำ
• Most Sensitive Screening Test: Serum TSH เป็นการตรวจคัดกรองที่มีความไวสูงสุดในการประเมินการทำงานของต่อมไทรอยด์"""
ws.cell(87, 10, exp_87)

print('--- Fixing Row 88 (Ref 86) ---')
q_88 = 'ข้อใดต่อไปนี้เป็นอาการแสดงที่พบได้ในภาวะต่อมไทรอยด์ทำงานเกิน (Hyperthyroidism)?'
ws.cell(88, 2, q_88)
ws.cell(88, 9, 1)
exp_88 = """✅ คำตอบที่ถูกต้อง: ข้อ ก. หงุดหงิดง่าย

💡 Background:
ฮอร์โมนไทรอยด์ทำหน้าที่เพิ่มอัตราการเผาผลาญพื้นฐานของร่างกาย (Basal Metabolic Rate: BMR) และเพิ่มความไวของระบบประสาทซิมพาเทติกต่อสารสื่อประสาทกลุ่ม Catecholamines

🎯 ทำไมข้อนี้ถึงถูก:
หงุดหงิดง่าย (Irritability / Nervousness / Emotional lability) เป็นอาการแสดงเด่นทางระบบประสาทที่พบบ่อยในผู้ป่วย Hyperthyroidism ร่วมกับอาการกระวนกระวาย อยู่ไม่นิ่ง มือสั่น นอนไม่หลับ ใจสั่น เหงื่อออกมาก ขี้ร้อน ทนร้อนไม่ได้ และน้ำหนักลดแม้จะรับประทานอาหารมากขึ้น

🔍 ข้ออื่นผิดเพราะอะไร:
• ข้อ ข. (ผมร่วง): แม้พบได้ในไทรอยด์ทั้งสองชนิด แต่ผมแห้งหยาบและหลุดร่วงง่าย มักเป็นลักษณะเด่นของภาวะขาดไทรอยด์ (Hypothyroidism) ร่วมกับขนคิ้วด้านนอกบางลง (Loss of outer third of eyebrows)
• ข้อ ค. (ผิวแห้ง): ผิวแห้ง เย็น หยาบกร้าน (Dry, cool, coarse skin) เป็นอาการแสดงคลาสสิกของ Hypothyroidism (ตรงข้ามกับ Hyperthyroidism ที่ผิวจะอุ่น ชื้น เรียบเนียน)
• ข้อ ง. (หัวใจเต้นช้า): หัวใจเต้นช้า (Bradycardia) เป็นอาการแสดงของ Hypothyroidism จาก sympathetic tone ที่ลดลง ส่วน Hyperthyroidism จะมีหัวใจเต้นเร็ว (Sinus tachycardia, Palpitations)
• ข้อ จ. (ท้องผูก): ท้องผูก (Constipation) เกิดจากการบีบตัวของลำไส้ลดลงใน Hypothyroidism ส่วน Hyperthyroidism ลำไส้จะบีบตัวไวขึ้น ถ่ายอุจจาระบ่อยขึ้น (Hyperdefecation / Diarrhea)

📖 Guideline อ้างอิง:
2016 American Thyroid Association (ATA) Guidelines for Diagnosis and Management of Hyperthyroidism and Other Causes of Thyrotoxicosis

📌 จุดจำก่อนสอบ:
• Hyperthyroidism: BMR สูง = ขี้ร้อน เหงื่อออก ใจสั่น กินจุแต่น้ำหนักลด มือสั่น หงุดหงิดง่าย ถ่ายบ่อย
• Hypothyroidism: BMR ต่ำ = ขี้หนาว ผิวแห้ง ท้องผูก เฉื่อยชา หัวใจเต้นช้า น้ำหนักขึ้น หน้าบวม หนังตาบวม (Myxedema)"""
ws.cell(88, 10, exp_88)

print('--- Fixing Row 89 (Ref 87) ---')
q_89 = """[สถานการณ์: หญิงตั้งครรภ์ อายุ 28 ปี อายุครรภ์ 4 เดือน (Second trimester: สัปดาห์ที่ 16) ได้รับการวินิจฉัยว่าเป็น Graves' hyperthyroidism แพทย์สั่งจ่ายยา Propylthiouracil (PTU) 50 mg 1x3 pc ร่วมกับ Propranolol 10 mg 1x3 pc]
จากการสั่งใช้ยาดังกล่าว ปัญหาจากการใช้ยา (Drug-Related Problem: DRP) ที่ชัดเจนที่สุดตามแนวทางการบริบาลทางเภสัชกรรมคือข้อใด?"""
ws.cell(89, 2, q_89)
ws.cell(89, 4, 'จ่ายยาเพิ่มอีก 1 ตัว')
ws.cell(89, 5, 'มียาที่ไม่เหมาะสม (Inappropriate drug)')
ws.cell(89, 6, 'มียาที่ความถี่ในการบริหารไม่เหมาะสม')
ws.cell(89, 7, 'มียาที่ไม่มีข้อบ่งใช้')
ws.cell(89, 8, 'เกิดปฏิกิริยาระหว่างยาที่รุนแรง')
ws.cell(89, 9, 2)
exp_89 = """✅ คำตอบที่ถูกต้อง: ข้อ ข. มียาที่ไม่เหมาะสม (Inappropriate drug)

💡 Background:
การรักษา Graves' disease ในสตรีตั้งครรภ์ต้องคำนึงถึงความเสี่ยงต่อทารกในครรภ์ (Teratogenicity) และความเป็นพิษต่อตับของมารดา (Maternal hepatotoxicity) โดยแบ่งการจัดการตามไตรมาสของการตั้งครรภ์

🎯 ทำไมข้อนี้ถึงถูก:
มียาที่ไม่เหมาะสม (Inappropriate drug) ตามแนวทางของ American Thyroid Association (ATA) และ Endocrine Society:
1) ไตรมาสแรก (First trimester, อายุครรภ์ 1-12 สัปดาห์): แนะนำให้ใช้ Propylthiouracil (PTU) เป็นยาทางเลือกแรก เนื่องจาก Methimazole (MMI) สัมพันธ์กับความเสี่ยงเกิดความพิการแต่กำเนิด (MMI embryopathy: Choanal atresia, Esophageal atresia, Aplasia cutis)
2) ไตรมาสที่สองและสาม (Second & Third trimesters, อายุครรภ์ตั้งแต่ 13 สัปดาห์ขึ้นไป): แนะนำให้พิจารณาเปลี่ยนยา (Switch) จาก PTU มาเป็น Methimazole (MMI) เนื่องจาก PTU มีความเสี่ยงต่อการเกิดพิษตับรุนแรงและตับวายเฉียบพลันในมารดา (Severe hepatotoxicity / Fulminant liver failure)
ผู้ป่วยรายนี้ตั้งครรภ์ได้ 4 เดือน (ไตรมาสที่ 2) การยังคงให้ PTU จึงมีความเสี่ยงต่อภาวะตับวายของมารดา จึงเป็น DRP ประเภท ยาไม่เหมาะสม ควรเปลี่ยนเป็น MMI ขนาดยาต่ำสุดที่คุม FT4 ให้อยู่ระดับ upper normal limit

🔍 ข้ออื่นผิดเพราะอะไร:
• ข้อ ก. (จ่ายยาเพิ่มอีก 1 ตัว): ยาต้านไทรอยด์และ Beta-blocker ขนาดต่ำเพื่อคุมอาการใจสั่นครอบคลุมข้อบ่งใช้แล้ว ไม่มีความจำเป็นต้องเพิ่มยาตัวอื่น
• ข้อ ค. (มียาที่ความถี่ในการบริหารไม่เหมาะสม): PTU ปกติบริหารวันละ 2-3 ครั้ง (TID) เนื่องจากมี half-life สั้น ซึ่งความถี่สอดคล้องกับเภสัชจลนศาสตร์
• ข้อ ง. (มียาที่ไม่มีข้อบ่งใช้): ยาทั้งสองตัวมีข้อบ่งใช้ชัดเจน (PTU สำหรับต้านไทรอยด์, Propranolol สำหรับบรรเทาอาการ Adrenergic hyperstimulation)
• ข้อ จ. (เกิดปฏิกิริยาระหว่างยาที่รุนแรง): PTU และ Propranolol สามารถใช้ร่วมกันได้อย่างปลอดภัย ไม่มี clinically significant DDI ที่เป็นข้อห้าม

📖 Guideline อ้างอิง:
2017 Guidelines of the American Thyroid Association for the Diagnosis and Management of Thyroid Disease During Pregnancy and the Postpartum & 2018 ETA Guidelines

📌 จุดจำก่อนสอบ:
• 1st Trimester Pregnancy: เลือกใช้ PTU (เลี่ยง MMI embryopathy)
• 2nd - 3rd Trimester: เปลี่ยนเป็น MMI (เลี่ยง PTU maternal hepatotoxicity)
• Beta-blocker in Pregnancy: ใช้ Propranolol ขนาดต่ำระยะสั้นได้ หากอาการดีขึ้นควรหยุดยาเพื่อเลี่ยง Fetal bradycardia, IUGR และ Hypoglycemia"""
ws.cell(89, 10, exp_89)

print('--- Removing Duplicate/Corrupt Rows 90, 91, 92 ---')
ws.delete_rows(90, 3)

# Renumber column 1 for the whole sheet
for i in range(3, ws.max_row + 1):
    ws.cell(i, 1, i - 2)

wb.save('PLE CC QUIZ.xlsx')
print('Successfully fixed and cleaned 4. Endocrine in PLE CC QUIZ.xlsx!')
