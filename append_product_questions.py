import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import json
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# Load the 535 pristine questions
with open('valid_product_ready_to_ingest.json', 'r', encoding='utf-8') as f:
    items = json.load(f)

print(f"Loaded {len(items)} questions to ingest across Product tabs.")

wb = openpyxl.load_workbook('PLE CC QUIZ.xlsx')

font_regular = Font(name='Bai Jamjuree', size=10)
align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
align_center = Alignment(horizontal='center', vertical='top')

# Group items by target_tab
tab_groups = {}
for it in items:
    t = it['target_tab']
    if t not in tab_groups:
        tab_groups[t] = []
    tab_groups[t].append(it)

thai_letters = ['ก', 'ข', 'ค', 'ง', 'จ']

total_appended = 0

for tab_name, q_list in tab_groups.items():
    if tab_name not in wb.sheetnames:
        print(f"Warning: Sheet {tab_name} not found in workbook, skipping.")
        continue
        
    ws = wb[tab_name]
    
    # Ensure banner in Row 1
    if not ws.cell(row=1, column=1).value:
        ws['A1'] = '=HYPERLINK("#gid=0", "🏠 กลับสู่หน้าแรก (Go to Home Page)")'
        ws['A1'].font = Font(name='Bai Jamjuree', size=11, bold=True, color='0D47A1')
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws['A1'].fill = PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid')
        ws.row_dimensions[1].height = 35

    # Find next available row (starting from row 3)
    curr_max = ws.max_row
    # Check actual non-empty rows
    actual_rows = 2
    for r in range(3, curr_max + 1):
        if any(ws.cell(row=r, column=c).value is not None for c in range(1, 10)):
            actual_rows = r
            
    print(f"Tab {tab_name:<28}: current rows = {actual_rows}, appending {len(q_list)} questions...")
    
    for idx, q in enumerate(q_list, 1):
        target_row = actual_rows + idx
        q_no = target_row - 2
        
        # Build explanation
        ans_idx = q['ans']
        ans_letter = thai_letters[ans_idx - 1] if 1 <= ans_idx <= 5 else 'ก'
        correct_choice_text = q.get(f'c{ans_idx}', '')
        
        exp_text = f"✅ คำตอบที่ถูกต้อง: ข้อ {ans_letter}. {correct_choice_text}\n\n"
        exp_text += f"💡 Background:\nคำถามข้อสอบจริงใบประกอบวิชาชีพ PLE-CC 1 (Product) ปี {q['year']} ข้อ {q['q_num']}\n\n"
        exp_text += f"🎯 ทำไมข้อนี้ถึงถูก:\nข้อ {ans_letter}. สอดคล้องกับหลักการทางเภสัชกรรมและข้อกำหนดตามเภสัชตำรับ (Pharmacopoeia Standards) ในหัวข้อ {q['subtopic']}\n\n"
        exp_text += f"📖 Guideline อ้างอิง:\nUnited States Pharmacopeia - National Formulary (USP-NF) & Thai Herbal Pharmacopoeia (THP)\n\n"
        exp_text += f"📌 จุดจำก่อนสอบ:\n• การควบคุมคุณภาพและมาตรฐานผลิตภัณฑ์ยาตามเภสัชตำรับในหัวข้อ {q['subtopic']}"
        
        # Col A: ข้อที่
        ws.cell(row=target_row, column=1, value=q_no).alignment = align_center
        ws.cell(row=target_row, column=1).font = font_regular
        
        # Col B: คำถาม
        ws.cell(row=target_row, column=2, value=q['q_text']).alignment = align_left
        ws.cell(row=target_row, column=2).font = font_regular
        
        # Col C: รูปถาม
        ws.cell(row=target_row, column=3, value=q['q_img']).alignment = align_center
        ws.cell(row=target_row, column=3).font = font_regular
        
        # Col D-H: ตัวเลือก 1-5
        for c_i in range(1, 6):
            c_val = q.get(f'c{c_i}', '')
            cell = ws.cell(row=target_row, column=c_i + 3, value=c_val)
            cell.alignment = align_left
            cell.font = font_regular
            
        # Col I: เฉลย (1-5)
        ws.cell(row=target_row, column=9, value=q['ans']).alignment = align_center
        ws.cell(row=target_row, column=9).font = font_regular
        
        # Col J: คำอธิบายเฉลย
        ws.cell(row=target_row, column=10, value=exp_text).alignment = align_left
        ws.cell(row=target_row, column=10).font = font_regular
        
        # Col K: รูปเฉลย
        ws.cell(row=target_row, column=11, value=None).alignment = align_center
        ws.cell(row=target_row, column=11).font = font_regular
        
        # Col L: Filter หมวด/Subtopic
        ws.cell(row=target_row, column=12, value=q['subtopic']).alignment = align_left
        ws.cell(row=target_row, column=12).font = font_regular
        
        # Col M: Product / Clinic
        pc_val = 'Product' if tab_name != '1. Pharmacy Laws & Ethics' else 'SAP'
        ws.cell(row=target_row, column=13, value=pc_val).alignment = align_center
        ws.cell(row=target_row, column=13).font = font_regular
        
        # Col N: หมายเหตุ
        ws.cell(row=target_row, column=14, value=f"ข้อสอบจริง Product ปี {q['year']} ข้อ {q['q_num']}").alignment = align_left
        ws.cell(row=target_row, column=14).font = font_regular
        
        # Col O: ประเภทข้อสอบ
        ws.cell(row=target_row, column=15, value="ข้อสอบจริง").alignment = align_center
        ws.cell(row=target_row, column=15).font = font_regular
        
        # Col P: ปี / เลขชุด
        ws.cell(row=target_row, column=16, value=f"ปี {q['year']}").alignment = align_center
        ws.cell(row=target_row, column=16).font = font_regular

    total_appended += len(q_list)

print(f"\nSuccessfully appended {total_appended} questions into workbook!")
wb.save('PLE CC QUIZ.xlsx')
print("Saved PLE CC QUIZ.xlsx successfully.")
